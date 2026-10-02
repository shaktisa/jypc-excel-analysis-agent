"""Deterministic Excel engine.

Every spreadsheet capability lives here as a pure, synchronously-testable function.
The LLM never manipulates a workbook directly - it may only choose one of these
functions and supply arguments, which keeps behaviour auditable and lets the whole
feature surface be covered by unit tests without calling a model.

Supported inputs: .xlsx, .xlsm, .csv, .tsv
"""
from __future__ import annotations

import csv
import io
import re
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

SPREADSHEET_SUFFIXES = {".xlsx", ".xlsm", ".csv", ".tsv"}
MAX_PREVIEW_ROWS = 200
MAX_PREVIEW_COLS = 40


class ExcelToolError(ValueError):
    """Raised for user-correctable problems (bad sheet name, bad range, ...)."""


# --------------------------------------------------------------------------- #
# Load / save
# --------------------------------------------------------------------------- #
def _delimited_to_workbook(data: bytes, delimiter: str, title: str) -> Workbook:
    text = data.decode("utf-8-sig", errors="replace")
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        ws.append([_coerce_scalar(cell) for cell in row])
    return wb


def _coerce_scalar(value: str) -> Any:
    """Turn CSV text into a number/bool where unambiguous, so stats work."""
    raw = value.strip()
    if raw == "":
        return None
    low = raw.lower()
    if low in {"true", "false"}:
        return low == "true"
    try:
        if re.fullmatch(r"[+-]?\d+", raw):
            return int(raw)
        return float(raw)
    except ValueError:
        return value


def load_workbook_from_bytes(data: bytes, filename: str) -> Workbook:
    """Load any supported spreadsheet format into an openpyxl Workbook."""
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SPREADSHEET_SUFFIXES:
        raise ExcelToolError(
            f"Unsupported file type '{suffix or filename}'. "
            f"Supported: {', '.join(sorted(SPREADSHEET_SUFFIXES))}"
        )
    if suffix == ".csv":
        return _delimited_to_workbook(data, ",", "Sheet1")
    if suffix == ".tsv":
        return _delimited_to_workbook(data, "\t", "Sheet1")
    try:
        return load_workbook(io.BytesIO(data), keep_vba=suffix == ".xlsm")
    except Exception as exc:  # corrupt upload
        raise ExcelToolError(f"Could not open workbook: {exc}") from exc


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def new_workbook(sheet_name: str = "Sheet1") -> Workbook:
    wb = Workbook()
    wb.active.title = sheet_name
    return wb


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _sheet(wb: Workbook, name: str | None) -> Worksheet:
    if not name:
        return wb.active
    if name not in wb.sheetnames:
        raise ExcelToolError(f"Sheet '{name}' not found. Available: {', '.join(wb.sheetnames)}")
    return wb[name]


def _bounds(ws: Worksheet, cell_range: str | None):
    if not cell_range:
        return 1, 1, max(ws.max_row, 1), max(ws.max_column, 1)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range.replace("$", ""))
    except Exception as exc:
        raise ExcelToolError(f"Invalid range '{cell_range}': {exc}") from exc
    return (
        min_row or 1,
        min_col or 1,
        max_row or ws.max_row or 1,
        max_col or ws.max_column or 1,
    )


def _hex(color: str | None) -> str | None:
    if not color:
        return None
    value = color.strip().lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}|[0-9A-F]{8}", value):
        raise ExcelToolError(f"Invalid colour '{color}'. Use hex like 'FFC000' or '#FFC000'.")
    return value if len(value) == 8 else f"FF{value}"


def _cell_value(cell) -> Any:
    value = cell.value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _as_dataframe(ws: Worksheet, cell_range: str | None = None) -> pd.DataFrame:
    """Read a region into a DataFrame, treating the first row as the header."""
    min_row, min_col, max_row, max_col = _bounds(ws, cell_range)
    rows = [
        [_cell_value(c) for c in row]
        for row in ws.iter_rows(
            min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col
        )
    ]
    rows = [r for r in rows if any(v is not None and str(v).strip() != "" for v in r)]
    if not rows:
        return pd.DataFrame()
    header, body = rows[0], rows[1:]
    header_is_text = all(isinstance(h, str) and h.strip() for h in header)
    if header_is_text and body:
        return pd.DataFrame(body, columns=[str(h).strip() for h in header])
    cols = [get_column_letter(min_col + i) for i in range(len(header))]
    return pd.DataFrame(rows, columns=cols)


# --------------------------------------------------------------------------- #
# Read / analyse
# --------------------------------------------------------------------------- #
def list_sheets(wb: Workbook) -> dict:
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        sheets.append(
            {
                "name": name,
                "rows": ws.max_row,
                "columns": ws.max_column,
                "dimensions": ws.dimensions,
                "charts": len(getattr(ws, "_charts", [])),
            }
        )
    return {"sheets": sheets, "active": wb.active.title}


def read_range(
    wb: Workbook, sheet: str | None = None, cell_range: str | None = None, max_rows: int = 50
) -> dict:
    """Return cell values (and formulas where present) for a region."""
    ws = _sheet(wb, sheet)
    min_row, min_col, max_row, max_col = _bounds(ws, cell_range)
    max_row = min(max_row, min_row + max(max_rows, 1) - 1)
    max_col = min(max_col, min_col + MAX_PREVIEW_COLS - 1)
    values = [
        [_cell_value(c) for c in row]
        for row in ws.iter_rows(min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col)
    ]
    return {
        "sheet": ws.title,
        "range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
        "columns": [get_column_letter(min_col + i) for i in range(max_col - min_col + 1)],
        "start_row": min_row,
        "values": values,
        "truncated": ws.max_row > max_row,
    }


def describe_sheet(wb: Workbook, sheet: str | None = None, cell_range: str | None = None) -> dict:
    """Profile a region: dtypes, null counts, numeric stats, top categories."""
    ws = _sheet(wb, sheet)
    df = _as_dataframe(ws, cell_range)
    if df.empty:
        return {"sheet": ws.title, "row_count": 0, "columns": []}

    columns: list[dict] = []
    for name in df.columns:
        series = df[name]
        numeric = pd.to_numeric(series, errors="coerce")
        is_numeric = numeric.notna().sum() >= max(1, int(0.8 * series.notna().sum()))
        info: dict[str, Any] = {
            "name": str(name),
            "inferred_type": "numeric" if is_numeric else "text",
            "non_null": int(series.notna().sum()),
            "nulls": int(series.isna().sum()),
            "distinct": int(series.nunique(dropna=True)),
        }
        if is_numeric and numeric.notna().any():
            info["stats"] = {
                "min": float(numeric.min()),
                "max": float(numeric.max()),
                "mean": round(float(numeric.mean()), 4),
                "median": float(numeric.median()),
                "sum": round(float(numeric.sum()), 4),
            }
        else:
            top = series.dropna().astype(str).value_counts().head(5)
            info["top_values"] = [{"value": k, "count": int(v)} for k, v in top.items()]
        columns.append(info)
    return {
        "sheet": ws.title,
        "row_count": int(len(df)),
        "column_count": int(len(df.columns)),
        "columns": columns,
    }


def compute_aggregate(
    wb: Workbook, sheet: str | None, cell_range: str, operation: str = "sum"
) -> dict:
    """Evaluate an aggregate over a range.

    openpyxl does not run Excel's calculation engine, so this gives the agent real
    numbers to reason about instead of an uncomputed formula string.
    """
    ws = _sheet(wb, sheet)
    min_row, min_col, max_row, max_col = _bounds(ws, cell_range)
    values = pd.to_numeric(
        pd.Series(
            [
                c.value
                for row in ws.iter_rows(
                    min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col
                )
                for c in row
            ]
        ),
        errors="coerce",
    ).dropna()
    ops = {
        "sum": values.sum,
        "mean": values.mean,
        "average": values.mean,
        "min": values.min,
        "max": values.max,
        "count": lambda: float(values.count()),
        "median": values.median,
        "std": values.std,
    }
    key = operation.strip().lower()
    if key not in ops:
        raise ExcelToolError(f"Unknown operation '{operation}'. Try: {', '.join(sorted(ops))}")
    if values.empty:
        return {"sheet": ws.title, "range": cell_range, "operation": key, "result": None}
    result = float(ops[key]())
    return {
        "sheet": ws.title,
        "range": cell_range,
        "operation": key,
        "result": round(result, 6),
        "values_counted": int(values.count()),
    }


# --------------------------------------------------------------------------- #
# Create / modify
# --------------------------------------------------------------------------- #
def add_sheet(wb: Workbook, name: str, index: int | None = None) -> dict:
    if name in wb.sheetnames:
        raise ExcelToolError(f"Sheet '{name}' already exists.")
    wb.create_sheet(title=name, index=index)
    return {"created": name, "sheets": wb.sheetnames}


def delete_sheet(wb: Workbook, name: str) -> dict:
    ws = _sheet(wb, name)
    if len(wb.sheetnames) == 1:
        raise ExcelToolError("Cannot delete the only sheet in a workbook.")
    wb.remove(ws)
    return {"deleted": name, "sheets": wb.sheetnames}


def rename_sheet(wb: Workbook, name: str, new_name: str) -> dict:
    ws = _sheet(wb, name)
    if new_name in wb.sheetnames:
        raise ExcelToolError(f"Sheet '{new_name}' already exists.")
    ws.title = new_name
    return {"renamed_from": name, "renamed_to": new_name, "sheets": wb.sheetnames}


def write_cells(
    wb: Workbook, sheet: str | None, anchor: str, values: list[list[Any]]
) -> dict:
    """Write a 2-D block of values with its top-left corner at ``anchor``."""
    ws = _sheet(wb, sheet)
    if not isinstance(values, list) or not values:
        raise ExcelToolError("'values' must be a non-empty list of rows.")
    start_row, start_col, *_ = _bounds(ws, anchor)
    written = 0
    for r, row in enumerate(values):
        row = row if isinstance(row, list) else [row]
        for c, value in enumerate(row):
            ws.cell(row=start_row + r, column=start_col + c, value=value)
            written += 1
    end = f"{get_column_letter(start_col + max(len(r) if isinstance(r, list) else 1 for r in values) - 1)}{start_row + len(values) - 1}"
    return {"sheet": ws.title, "range": f"{anchor}:{end}", "cells_written": written}


def set_formula(wb: Workbook, sheet: str | None, cell: str, formula: str) -> dict:
    """Set an Excel formula. Excel/LibreOffice evaluates it when the file opens."""
    ws = _sheet(wb, sheet)
    text = formula if formula.startswith("=") else f"={formula}"
    ws[cell.replace("$", "")] = text
    return {
        "sheet": ws.title,
        "cell": cell,
        "formula": text,
        "note": "Formulas are evaluated by Excel on open; use compute_aggregate for a value now.",
    }


def fill_formula(
    wb: Workbook, sheet: str | None, cell_range: str, formula_template: str
) -> dict:
    """Fill a column/row with a formula, substituting ``{row}`` and ``{col}``."""
    ws = _sheet(wb, sheet)
    min_row, min_col, max_row, max_col = _bounds(ws, cell_range)
    count = 0
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            text = formula_template.replace("{row}", str(r)).replace("{col}", get_column_letter(c))
            ws.cell(row=r, column=c, value=text if text.startswith("=") else f"={text}")
            count += 1
    return {"sheet": ws.title, "range": cell_range, "cells_filled": count}


def sort_range(
    wb: Workbook, sheet: str | None, cell_range: str, by_column: str, descending: bool = False
) -> dict:
    """Sort a region in place, keeping the header row fixed."""
    ws = _sheet(wb, sheet)
    min_row, min_col, max_row, max_col = _bounds(ws, cell_range)
    header = [ws.cell(row=min_row, column=c).value for c in range(min_col, max_col + 1)]
    try:
        offset = [str(h).strip() for h in header].index(str(by_column).strip())
    except ValueError:
        if len(by_column) <= 2 and by_column.isalpha():
            offset = range_boundaries(f"{by_column}1")[0] - min_col
        else:
            raise ExcelToolError(
                f"Column '{by_column}' not in header {header}."
            ) from None
    body = [
        [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
        for r in range(min_row + 1, max_row + 1)
    ]
    body.sort(key=lambda row: (row[offset] is None, _sort_key(row[offset])), reverse=descending)
    for i, row in enumerate(body):
        for j, value in enumerate(row):
            ws.cell(row=min_row + 1 + i, column=min_col + j, value=value)
    return {"sheet": ws.title, "range": cell_range, "sorted_by": by_column, "rows": len(body)}


def _sort_key(value: Any):
    if value is None:
        return (1, 0.0, "")
    if isinstance(value, int | float) and not isinstance(value, bool):
        return (0, float(value), "")
    return (0, 0.0, str(value).lower())


# --------------------------------------------------------------------------- #
# Formatting
# --------------------------------------------------------------------------- #
def format_cells(
    wb: Workbook,
    sheet: str | None,
    cell_range: str,
    bold: bool | None = None,
    italic: bool | None = None,
    font_size: float | None = None,
    font_color: str | None = None,
    fill_color: str | None = None,
    number_format: str | None = None,
    align: str | None = None,
    wrap_text: bool | None = None,
    border: bool | None = None,
    column_width: float | None = None,
) -> dict:
    """Apply font, fill, number-format, alignment and border styling to a range."""
    ws = _sheet(wb, sheet)
    min_row, min_col, max_row, max_col = _bounds(ws, cell_range)
    fill = (
        PatternFill(start_color=_hex(fill_color), end_color=_hex(fill_color), fill_type="solid")
        if fill_color
        else None
    )
    side = Side(style="thin", color="FF999999")
    applied = 0
    for row in ws.iter_rows(min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col):
        for cell in row:
            font = cell.font
            cell.font = Font(
                name=font.name,
                size=font_size if font_size is not None else font.size,
                bold=bold if bold is not None else font.bold,
                italic=italic if italic is not None else font.italic,
                color=_hex(font_color) if font_color else font.color,
            )
            if fill:
                cell.fill = fill
            if number_format:
                cell.number_format = number_format
            if align or wrap_text is not None:
                cell.alignment = Alignment(
                    horizontal=align or cell.alignment.horizontal,
                    vertical=cell.alignment.vertical or "center",
                    wrap_text=wrap_text if wrap_text is not None else cell.alignment.wrap_text,
                )
            if border:
                cell.border = Border(left=side, right=side, top=side, bottom=side)
            applied += 1
    if column_width:
        for c in range(min_col, max_col + 1):
            ws.column_dimensions[get_column_letter(c)].width = column_width
    return {"sheet": ws.title, "range": cell_range, "cells_formatted": applied}


def autofit_columns(wb: Workbook, sheet: str | None = None, max_width: float = 60) -> dict:
    """Approximate Excel's autofit by measuring the longest value per column."""
    ws = _sheet(wb, sheet)
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), max_width)
    return {"sheet": ws.title, "columns_resized": len(widths)}


def conditional_format(
    wb: Workbook,
    sheet: str | None,
    cell_range: str,
    rule_type: str,
    operator: str | None = None,
    value: Any = None,
    value2: Any = None,
    formula: str | None = None,
    color: str = "FFC7CE",
) -> dict:
    """Add a conditional-formatting rule.

    ``rule_type`` is one of: ``color_scale``, ``data_bar``, ``cell_is``, ``formula``.
    """
    ws = _sheet(wb, sheet)
    _bounds(ws, cell_range)  # validate
    kind = rule_type.strip().lower()
    if kind == "color_scale":
        rule = ColorScaleRule(
            start_type="min", start_color=_hex("F8696B"),
            mid_type="percentile", mid_value=50, mid_color=_hex("FFEB84"),
            end_type="max", end_color=_hex("63BE7B"),
        )
    elif kind == "data_bar":
        rule = DataBarRule(start_type="min", end_type="max", color=_hex(color) or "FF638EC6")
    elif kind == "cell_is":
        if not operator:
            raise ExcelToolError("'operator' is required for a cell_is rule (e.g. greaterThan).")
        operands = [str(v) for v in (value, value2) if v is not None]
        rule = CellIsRule(
            operator=operator,
            formula=operands or ["0"],
            fill=PatternFill(start_color=_hex(color), end_color=_hex(color), fill_type="solid"),
        )
    elif kind == "formula":
        if not formula:
            raise ExcelToolError("'formula' is required for a formula rule.")
        rule = FormulaRule(
            formula=[formula.lstrip("=")],
            fill=PatternFill(start_color=_hex(color), end_color=_hex(color), fill_type="solid"),
        )
    else:
        raise ExcelToolError(
            f"Unknown rule_type '{rule_type}'. Use color_scale, data_bar, cell_is or formula."
        )
    ws.conditional_formatting.add(cell_range, rule)
    return {"sheet": ws.title, "range": cell_range, "rule": kind}


# --------------------------------------------------------------------------- #
# Charts
# --------------------------------------------------------------------------- #
_CHART_TYPES = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart,
                "area": AreaChart, "scatter": ScatterChart}


def create_chart(
    wb: Workbook,
    sheet: str | None,
    chart_type: str,
    data_range: str,
    categories_range: str | None = None,
    title: str | None = None,
    anchor: str = "H2",
    target_sheet: str | None = None,
    titles_from_data: bool = True,
) -> dict:
    """Create a native Excel chart from a data range and anchor it on a sheet."""
    ws = _sheet(wb, sheet)
    kind = chart_type.strip().lower()
    factory = _CHART_TYPES.get(kind)
    if not factory:
        raise ExcelToolError(
            f"Unknown chart_type '{chart_type}'. Use: {', '.join(sorted(_CHART_TYPES))}"
        )
    chart = factory()
    if kind == "column":
        chart.type = "col"
    elif kind == "bar":
        chart.type = "bar"
    if title:
        chart.title = title

    d_min_row, d_min_col, d_max_row, d_max_col = _bounds(ws, data_range)
    data_ref = Reference(
        ws, min_col=d_min_col, min_row=d_min_row, max_col=d_max_col, max_row=d_max_row
    )
    if kind == "scatter":
        if not categories_range:
            raise ExcelToolError("A scatter chart needs 'categories_range' for the X values.")
        x_min_row, x_min_col, x_max_row, x_max_col = _bounds(ws, categories_range)
        xvalues = Reference(
            ws, min_col=x_min_col, min_row=x_min_row, max_col=x_max_col, max_row=x_max_row
        )
        for col in range(d_min_col, d_max_col + 1):
            values = Reference(ws, min_col=col, min_row=d_min_row, max_row=d_max_row)
            chart.series.append(Series(values, xvalues, title_from_data=titles_from_data))
    else:
        chart.add_data(data_ref, titles_from_data=titles_from_data)
        if categories_range:
            c_min_row, c_min_col, c_max_row, c_max_col = _bounds(ws, categories_range)
            chart.set_categories(
                Reference(
                    ws, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row
                )
            )
    chart.height, chart.width = 8, 16
    destination = _sheet(wb, target_sheet) if target_sheet else ws
    destination.add_chart(chart, anchor)
    return {
        "sheet": destination.title,
        "chart_type": kind,
        "data_range": data_range,
        "anchor": anchor,
        "title": title,
    }


# --------------------------------------------------------------------------- #
# UI preview
# --------------------------------------------------------------------------- #
def preview(wb: Workbook, sheet: str | None = None, max_rows: int = MAX_PREVIEW_ROWS) -> dict:
    """Compact payload the frontend renders as a grid."""
    ws = _sheet(wb, sheet)
    data = read_range(wb, ws.title, None, max_rows=max_rows)
    data["sheets"] = wb.sheetnames
    return data
