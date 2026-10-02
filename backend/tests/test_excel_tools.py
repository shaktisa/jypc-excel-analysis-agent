"""Unit tests for the deterministic Excel engine (P0 read/analyse/chart surface)."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from app.excel_tools import (
    ExcelToolError,
    compute_aggregate,
    create_chart,
    describe_sheet,
    list_sheets,
    load_workbook_from_bytes,
    preview,
    read_range,
    workbook_to_bytes,
)

# --------------------------------------------------------------------------- #
# Loading every supported format
# --------------------------------------------------------------------------- #


def test_loads_xlsx_roundtrip(sales_bytes):
    wb = load_workbook_from_bytes(sales_bytes, "sales.xlsx")
    assert wb.sheetnames == ["Sales"]
    assert wb["Sales"].cell(row=2, column=1).value == "North"


def test_loads_csv_and_coerces_numbers():
    data = b"Name,Score,Active\nAda,91,true\nGrace,88,false\n"
    wb = load_workbook_from_bytes(data, "scores.csv")
    ws = wb.active
    assert ws.cell(row=2, column=2).value == 91
    assert ws.cell(row=2, column=2).value.__class__ is int
    assert ws.cell(row=2, column=3).value is True


def test_loads_tsv():
    wb = load_workbook_from_bytes(b"A\tB\n1\t2\n", "data.tsv")
    assert wb.active.cell(row=2, column=2).value == 2


def test_loads_csv_with_bom_and_blanks():
    wb = load_workbook_from_bytes("\ufeffName,Score\nAda,\n".encode(), "s.csv")
    assert wb.active.cell(row=1, column=1).value == "Name"
    assert wb.active.cell(row=2, column=2).value is None


def test_rejects_unsupported_extension():
    with pytest.raises(ExcelToolError, match="Unsupported file type"):
        load_workbook_from_bytes(b"x", "notes.pdf")


def test_rejects_corrupt_workbook():
    with pytest.raises(ExcelToolError, match="Could not open workbook"):
        load_workbook_from_bytes(b"definitely not a zip", "broken.xlsx")


# --------------------------------------------------------------------------- #
# Structure + reading
# --------------------------------------------------------------------------- #


def test_list_sheets_reports_shape(sales_workbook):
    result = list_sheets(sales_workbook)
    assert result["active"] == "Sales"
    assert result["sheets"][0]["rows"] == 9
    assert result["sheets"][0]["columns"] == 4


def test_read_range_whole_sheet(sales_workbook):
    result = read_range(sales_workbook, "Sales")
    assert result["values"][0] == ["Region", "Quarter", "Revenue", "Units"]
    assert len(result["values"]) == 9


def test_read_range_subset_and_truncation(sales_workbook):
    result = read_range(sales_workbook, "Sales", "A1:B3")
    assert result["values"] == [["Region", "Quarter"], ["North", "Q1"], ["North", "Q2"]]
    assert result["range"] == "A1:B3"

    limited = read_range(sales_workbook, "Sales", None, max_rows=3)
    assert len(limited["values"]) == 3
    assert limited["truncated"] is True


def test_unknown_sheet_raises(sales_workbook):
    with pytest.raises(ExcelToolError, match="not found"):
        read_range(sales_workbook, "Nope")


def test_invalid_range_raises(sales_workbook):
    with pytest.raises(ExcelToolError, match="Invalid range"):
        read_range(sales_workbook, "Sales", "not-a-range")


def test_preview_includes_sheet_list(sales_workbook):
    result = preview(sales_workbook)
    assert result["sheets"] == ["Sales"]


# --------------------------------------------------------------------------- #
# Profiling / analysis
# --------------------------------------------------------------------------- #


def test_describe_sheet_profiles_columns(sales_workbook):
    profile = describe_sheet(sales_workbook, "Sales")
    assert profile["row_count"] == 8
    by_name = {c["name"]: c for c in profile["columns"]}

    region = by_name["Region"]
    assert region["inferred_type"] == "text"
    assert region["distinct"] == 4
    assert {t["value"] for t in region["top_values"]} == {"North", "South", "East", "West"}

    revenue = by_name["Revenue"]
    assert revenue["inferred_type"] == "numeric"
    assert revenue["nulls"] == 1, "the deliberate blank must be surfaced"
    assert revenue["stats"]["max"] == 232900
    assert revenue["stats"]["min"] == 76300


def test_describe_empty_sheet_is_safe():
    from app.excel_tools import new_workbook

    profile = describe_sheet(new_workbook("Blank"))
    assert profile["row_count"] == 0
    assert profile["columns"] == []


@pytest.mark.parametrize(
    ("operation", "expected"),
    [
        ("sum", 973350.0),
        ("min", 76300.0),
        ("max", 232900.0),
        ("count", 7.0),
    ],
)
def test_compute_aggregate(sales_workbook, operation, expected):
    result = compute_aggregate(sales_workbook, "Sales", "C2:C9", operation)
    assert result["result"] == pytest.approx(expected)


def test_compute_aggregate_ignores_nulls(sales_workbook):
    assert compute_aggregate(sales_workbook, "Sales", "C2:C9", "count")["values_counted"] == 7


def test_compute_aggregate_rejects_unknown_operation(sales_workbook):
    with pytest.raises(ExcelToolError, match="Unknown operation"):
        compute_aggregate(sales_workbook, "Sales", "C2:C9", "kurtosis")


def test_compute_aggregate_on_text_returns_none(sales_workbook):
    assert compute_aggregate(sales_workbook, "Sales", "A2:A9", "sum")["result"] is None


# --------------------------------------------------------------------------- #
# Charts
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize("chart_type", ["bar", "column", "line", "pie", "area"])
def test_create_chart_persists_into_workbook(sales_workbook, chart_type):
    result = create_chart(
        sales_workbook, "Sales", chart_type, "C1:C9", "A2:A9", f"{chart_type} test"
    )
    assert result["chart_type"] == chart_type
    reloaded = load_workbook(__import__("io").BytesIO(workbook_to_bytes(sales_workbook)))
    assert len(reloaded["Sales"]._charts) == 1, "chart must survive a save/load round-trip"


def test_scatter_chart_requires_categories(sales_workbook):
    with pytest.raises(ExcelToolError, match="scatter chart needs"):
        create_chart(sales_workbook, "Sales", "scatter", "C1:C9")


def test_create_chart_rejects_unknown_type(sales_workbook):
    with pytest.raises(ExcelToolError, match="Unknown chart_type"):
        create_chart(sales_workbook, "Sales", "sunburst", "C1:C9")


def test_chart_can_target_another_sheet(sales_workbook):
    sales_workbook.create_sheet("Dashboard")
    result = create_chart(
        sales_workbook, "Sales", "column", "C1:C9", "A2:A9", "Revenue", "B2", "Dashboard"
    )
    assert result["sheet"] == "Dashboard"
    assert len(sales_workbook["Dashboard"]._charts) == 1
